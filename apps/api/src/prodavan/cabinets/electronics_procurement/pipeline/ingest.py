"""Honest ingest: csv/txt always; xlsx via table filter; .xls stays not_implemented."""

from __future__ import annotations

import csv
import io
from pathlib import Path

from prodavan.cabinets.electronics_procurement.pipeline.spec_table import (
    clean_extracted_rows,
    headers_by_sheet,
    to_ingest_rows,
)


def ingest_rows(input_path: Path) -> dict:
    suffix = input_path.suffix.lower()
    if suffix == ".csv":
        return {"parser": "csv", "rows": _parse_csv(input_path.read_text(encoding="utf-8-sig"))}
    if suffix == ".txt":
        return {"parser": "txt", "rows": _parse_txt(input_path.read_text(encoding="utf-8-sig"))}
    if suffix == ".xlsx":
        return _parse_xlsx(input_path)
    return {
        "parser": "not_implemented",
        "parser_status": "not_implemented",
        "note": f"No parser for {suffix or 'unknown'}; rows not invented",
        "rows": [],
    }


def extracted_markdown(filename: str, ingest: dict) -> str:
    if ingest.get("parser_status") == "not_implemented":
        return (
            f"# {filename}\n\n"
            f"parser_status: not_implemented\n\n"
            "Binary spec not extracted. Rows were not invented.\n"
        )
    if ingest.get("parser_status") == "error":
        return (
            f"# {filename}\n\n"
            f"parser_status: error\n\n"
            f"{ingest.get('note', 'Workbook unreadable')}\n\n"
            "Rows were not invented.\n"
        )
    lines = [f"# {filename}", "", f"parser: {ingest.get('parser')}", ""]
    for row in ingest["rows"]:
        sheet = row.get("sheet")
        prefix = f"[{sheet}] " if sheet else ""
        lines.append(f"- {prefix}{row.get('raw_line', '')}")
    return "\n".join(lines) + "\n"


def _parse_csv(text: str) -> list[dict]:
    reader = csv.DictReader(io.StringIO(text))
    rows: list[dict] = []
    if reader.fieldnames is None:
        return rows
    for index, cells in enumerate(reader, start=2):
        cleaned = {k: (v or "").strip() for k, v in cells.items() if k}
        raw = " ".join(v for v in cleaned.values() if v)
        if not raw:
            continue
        rows.append({"row_index": index, "cells": cleaned, "raw_line": raw})
    return rows


def _parse_txt(text: str) -> list[dict]:
    rows: list[dict] = []
    for index, line in enumerate(text.splitlines(), start=1):
        raw = line.strip()
        if not raw:
            continue
        rows.append({"row_index": index, "cells": {"line": raw}, "raw_line": raw})
    return rows


def _parse_xlsx(path: Path) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {
            "parser": "xlsx",
            "parser_status": "error",
            "note": "openpyxl not installed; rows not invented",
            "rows": [],
        }
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — corrupt office zip is expected
        return {
            "parser": "xlsx",
            "parser_status": "error",
            "note": f"unreadable xlsx: {exc}; rows not invented",
            "rows": [],
        }
    raw_rows: list[dict] = []
    try:
        for sheet_name in workbook.sheetnames:
            if sheet_name.startswith("_"):
                continue
            worksheet = workbook[sheet_name]
            for index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                cells = ["" if cell is None else str(cell).strip() for cell in row]
                if not any(cells):
                    continue
                raw_rows.append(
                    {
                        "sheet": sheet_name,
                        "row": index,
                        "cells": cells,
                        "raw": " | ".join(c for c in cells if c),
                    }
                )
    finally:
        workbook.close()

    header_map = headers_by_sheet(raw_rows)
    cleaned = clean_extracted_rows(raw_rows)
    return {"parser": "xlsx", "rows": to_ingest_rows(cleaned, header_map)}
