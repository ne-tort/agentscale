"""Мастер-прайс «Подбора техники»: виртуальная таблица из OpenSearch.

Данные НЕ хранятся в БД модуля: строки читаются из OS-индексов готовых
каталогов кабинета по списку поставщиков с флагом ``master_price``
(trusted_sellers). UI-страница ходит постранично (from/size), экспорт —
курсором search_after батчами (сотни тысяч строк, from/size упирается в
max_result_window).

Колонки и лист повторяют мастер-прайс легаси-Commerce
(s4b_catalog.masterprice): лист «Прайс», шапка в строке 1, данные со
строки 2, freeze A2, колонки A–K.
"""

from __future__ import annotations

import io
from collections.abc import AsyncIterator
from typing import Any

from sqlalchemy.ext.asyncio import AsyncSession

from prodavan.application.modules.equipment_catalog_opensearch import (
    OS_NAMESPACE,
    catalog_os_index_name,
)
from prodavan.domain.search_index.types import MAX_SEARCH_SIZE

MASTER_PRICE_SHEET = "Прайс"

MASTER_PRICE_HEADERS = [
    "Поставщик",
    "Вид оборудования",
    "Бренд",
    "PN",
    "Наименование",
    "Есть",
    "Свободно",
    "Транзит",
    "Цена",
    "Валюта",
    "РРЦ",
]

MASTER_PRICE_COLUMN_WIDTHS = [20, 24, 16, 22, 60, 12, 12, 12, 14, 8, 14]

CURRENCY_DISPLAY = {"RUB": "руб", "USD": "usd", "EUR": "eur"}

# Стабильный сортируемый курсор: supplier/title — keyword-поля (title.raw),
# src_hash — уникальный тай-брейк для search_after.
MASTER_PRICE_SORT = [
    {"supplier": {"order": "asc"}},
    {"title.raw": {"order": "asc"}},
    {"src_hash": {"order": "asc"}},
]

_PAGE_SIZE_MAX = 200
_FROM_MAX = 10000 - _PAGE_SIZE_MAX
# SearchIndexService caps size at MAX_SEARCH_SIZE — батч должен равняться
# фактическому размеру страницы, иначе «неполный» батч оборвёт обход
# (инцидент: экспорт резался первыми 200 доками индекса).
_EXPORT_BATCH = min(1000, MAX_SEARCH_SIZE)
# страховка от бесконечного обхода (курсор должен двигаться):
# абсолютный потолок строк одного экспорта
_EXPORT_MAX_ROWS = 500_000


def master_price_sellers(seller_rows: list[dict[str, Any]]) -> list[str]:
    """Имена поставщиков, включённых в мастер-прайс (и не отключённых)."""
    out: list[str] = []
    for row in seller_rows:
        body = row.get("body") if isinstance(row.get("body"), dict) else row
        if not isinstance(body, dict):
            continue
        if body.get("master_price") is not True:
            continue
        if body.get("is_enabled") is False:
            continue
        name = str(body.get("name") or "").strip()
        if name and name not in out:
            out.append(name)
    return out


def build_master_price_query(suppliers: list[str], search: str) -> dict[str, Any]:
    """bool-filter по поставщикам + опциональный поиск по названию/P/N/бренду."""
    clauses: list[dict[str, Any]] = [{"terms": {"supplier": suppliers}}]
    term = (search or "").strip()
    if term:
        clauses.append(
            {
                "bool": {
                    "should": [
                        {"match": {"title": {"query": term, "operator": "and"}}},
                        {"term": {"part_number": term}},
                        {"term": {"part_number": term.upper()}},
                        {"term": {"part_number": term.lower()}},
                        {"term": {"brand": term}},
                    ],
                    "minimum_should_match": 1,
                }
            }
        )
    return {"bool": {"filter": clauses}}


def doc_to_master_row(doc: dict[str, Any]) -> dict[str, Any]:
    """OS-док → строка виртуальной таблицы (ключи = колонки мета-вьюхи)."""
    price = doc.get("price_num")
    rrc = doc.get("rrc_num")
    return {
        "supplier": str(doc.get("supplier") or ""),
        "category": "",
        "brand": str(doc.get("brand") or ""),
        "part_number": str(doc.get("part_number") or ""),
        "title": str(doc.get("title") or ""),
        "in_stock": bool(doc.get("in_stock")),
        "lead_time": str(doc.get("lead_time") or ""),
        "price": float(price) if isinstance(price, (int, float)) else None,
        "rrc": float(rrc) if isinstance(rrc, (int, float)) else None,
        "currency": str(doc.get("currency") or ""),
        "src_hash": str(doc.get("src_hash") or ""),
        "catalog_id": str(doc.get("catalog_id") or ""),
    }


def _qty_cell(in_stock: bool) -> str:
    # Количеств в каталоге нет: численное «Есть/Свободно» недоступно —
    # как в Commerce для нераспознанных количеств.
    return "По запросу"


def _price_cell(price: float | None) -> Any:
    return price if price is not None else "По запросу"


def _currency_cell(currency: str) -> str:
    code = (currency or "").strip().upper()
    return CURRENCY_DISPLAY.get(code, code.lower())


def master_price_row_values(row: dict[str, Any]) -> list[Any]:
    """Строка виртуальной таблицы → ячейки A–K шаблона Commerce."""
    qty = _qty_cell(bool(row.get("in_stock")))
    return [
        row.get("supplier") or "",
        row.get("category") or "",
        row.get("brand") or "",
        row.get("part_number") or "",
        row.get("title") or "",
        qty,
        qty,
        "В наличии" if row.get("in_stock") else "Под заказ",
        _price_cell(row.get("price")),
        _currency_cell(str(row.get("currency") or "")),
        _price_cell(row.get("rrc")),
    ]


async def try_list_virtual_page(
    session: Any,
    *,
    cabinet_id: str,
    table_slug: str,
    principal: Any,
    employee: Any,
    page: int,
    page_size: int,
    search: str,
) -> dict[str, Any] | None:
    """Виртуальные таблицы модуля (master_price): строки из OpenSearch.

    None — таблица обычная (DB); вызывающий идёт в instance-строки.
    """
    if table_slug != "master_price":
        return None
    return await MasterPriceService(
        session, principal=principal, employee=employee
    ).list_page(
        cabinet_id=cabinet_id, search=search, page=page, page_size=page_size
    )


def apply_margin(price: float | None, margin_pct: float) -> float | None:
    """Цена/РРЦ мастер-прайса с наценкой поставщика (Commerce: round 2)."""
    if price is None or not margin_pct:
        return price
    return round(float(price) * (1.0 + margin_pct / 100.0), 2)


def _apply_margin_to_row(row: dict, margins: dict) -> None:
    pct = margins.get(str(row.get("supplier") or "").strip().casefold(), 0.0)
    if not pct:
        return
    row["price"] = apply_margin(row.get("price"), pct)
    row["rrc"] = apply_margin(row.get("rrc"), pct)


class MasterPriceService:
    """Чтение мастер-прайса из OS-индексов готовых каталогов кабинета."""

    def __init__(
        self,
        session: AsyncSession,
        *,
        principal: Any = None,
        employee: Any = None,
    ) -> None:
        self._session = session
        self._principal = principal
        self._employee = employee

    async def _ready_indexes(self, cabinet_id: str) -> list[str]:
        rows = await self._module_rows(cabinet_id, "catalogs")
        out: list[str] = []
        for row in rows:
            raw_body = row.get("body")
            body: dict[str, Any] = raw_body if isinstance(raw_body, dict) else {}
            if str(body.get("status") or "").strip().lower() != "ready":
                continue
            out.append(catalog_os_index_name(str(row.get("row_id") or "")))
        return out

    async def _sellers(self, cabinet_id: str) -> list[str]:
        rows = await self._module_rows(cabinet_id, "trusted_sellers")
        return master_price_sellers([r for r in rows if isinstance(r, dict)])

    async def _margins(self, cabinet_id: str) -> dict[str, float]:
        """Наценка поставщика (margin_pct, %) для мастер-прайса.

        Как в Commerce: наценка применяется ТОЛЬКО при формировании
        мастер-прайса (страница/xlsx); в каталоге и офферах хранятся
        настоящие цены поставщика до наценки.
        """
        rows = await self._module_rows(cabinet_id, "trusted_sellers")
        out: dict[str, float] = {}
        for row in rows:
            body = row.get("body") if isinstance(row.get("body"), dict) else {}
            name = str(body.get("name") or "").strip().casefold()
            if not name:
                continue
            pct = body.get("margin_pct")
            if isinstance(pct, (int, float)):
                out[name] = float(pct)
        return out

    async def _module_rows(self, cabinet_id: str, table_slug: str) -> list[dict]:
        from prodavan.application.cabinets.cabinet_module_service import (
            CabinetModuleService,
        )

        svc = CabinetModuleService(self._session)
        return await svc.list_data_rows(
            cabinet_id=cabinet_id,
            module_id="mod_equipment",
            table_slug=table_slug,
            principal=self._principal,
            employee=self._employee,
        )

    async def list_page(
        self,
        *,
        cabinet_id: str,
        search: str = "",
        page: int = 1,
        page_size: int = 50,
    ) -> dict[str, Any]:
        from prodavan.core.infra.opensearch_manager import get_search_index_service

        size = max(1, min(int(page_size or 50), _PAGE_SIZE_MAX))
        page = max(1, int(page or 1))
        from_ = (page - 1) * size
        if from_ > _FROM_MAX:
            # глубокие страницы дешёвым способом недоступны: UI ограничивает
            return {"items": [], "total": 0, "page": page, "page_size": size}
        suppliers = await self._sellers(cabinet_id)
        if not suppliers:
            return {"items": [], "total": 0, "page": page, "page_size": size}
        indexes = await self._ready_indexes(cabinet_id)
        if not indexes:
            return {"items": [], "total": 0, "page": page, "page_size": size}
        svc = get_search_index_service()
        query = build_master_price_query(suppliers, search)
        margins = await self._margins(cabinet_id)
        items: list[dict[str, Any]] = []
        total = 0
        for index in indexes:
            result = await svc.search(
                namespace=OS_NAMESPACE,
                index=index,
                query=query,
                from_=from_,
                size=size,
                company_id="platform",
                cabinet_id=None,
                project_id=None,
                apply_tenant_filter=False,
                session=self._session,
                sort=MASTER_PRICE_SORT,
            )
            total += result.total
            for hit in result.hits:
                row = doc_to_master_row(hit.source)
                _apply_margin_to_row(row, margins)
                items.append(row)
        items.sort(key=lambda r: (r["supplier"], r["title"], r["src_hash"]))
        return {"items": items, "total": total, "page": page, "page_size": size}

    async def iter_rows(self, *, cabinet_id: str) -> AsyncIterator[dict[str, Any]]:
        """Полный обход курсором search_after (экспорт сотен тысяч строк)."""
        from prodavan.core.infra.opensearch_manager import get_search_index_service

        suppliers = await self._sellers(cabinet_id)
        if not suppliers:
            return
        indexes = await self._ready_indexes(cabinet_id)
        svc = get_search_index_service()
        query = build_master_price_query(suppliers, "")
        margins = await self._margins(cabinet_id)
        exported = 0
        for index in indexes:
            cursor: list[Any] | None = None
            while True:
                result = await svc.search(
                    namespace=OS_NAMESPACE,
                    index=index,
                    query=query,
                    from_=0,
                    size=_EXPORT_BATCH,
                    company_id="platform",
                    cabinet_id=None,
                    project_id=None,
                    apply_tenant_filter=False,
                    session=self._session,
                    sort=MASTER_PRICE_SORT,
                    search_after=cursor,
                )
                hits = list(result.hits)
                if not hits:
                    break
                next_cursor = list(hits[-1].sort or [])
                # курсор обязан двигаться: иначеOS вернёт ту же страницу и
                # обход станет бесконечным (инцидент: OOM в тесте-фейке)
                if not next_cursor or next_cursor == cursor:
                    for hit in hits:
                        if cursor is None or list(hit.sort or []) > cursor:
                            row = doc_to_master_row(hit.source)
                            _apply_margin_to_row(row, margins)
                            yield row
                            exported += 1
                    break
                for hit in hits:
                    row = doc_to_master_row(hit.source)
                    _apply_margin_to_row(row, margins)
                    yield row
                    exported += 1
                    if exported >= _EXPORT_MAX_ROWS:
                        return
                cursor = next_cursor
                if len(hits) < _EXPORT_BATCH:
                    break


def render_master_price_workbook(template_bytes: bytes, rows: list[dict[str, Any]]) -> bytes:
    """Заполняет шаблон (лист «Прайс», шапка в строке 1) строками мастер-прайса."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    sheet = wb[MASTER_PRICE_SHEET] if MASTER_PRICE_SHEET in wb.sheetnames else wb.active
    for idx, row in enumerate(rows):
        excel_row = idx + 2
        for col, value in enumerate(master_price_row_values(row), start=1):
            sheet.cell(row=excel_row, column=col, value=value)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
