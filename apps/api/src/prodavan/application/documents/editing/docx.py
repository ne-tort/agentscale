"""Word document editing helpers (python-docx, DOCUM).

Pure functions: bytes in → bytes out. ``set_placeholders`` rewrites ``{{key}}``
markers inside paragraphs and table cells with mapping values.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any

from docx import Document

from prodavan.domain.errors import AppError

_PLACEHOLDER_RE = re.compile(r"\{\{\s*([a-zA-Z0-9_.\-]+)\s*\}\}")


@dataclass(slots=True)
class ParagraphSpec:
    """Declarative paragraph for ``create_docx`` / ``append_paragraphs``."""

    text: str = ""
    style: str | None = None
    heading: int | None = None
    bold: bool = False


def _open(data: bytes) -> Document:
    if not data:
        raise AppError(
            code="VALIDATION_ERROR",
            title="Validation Error",
            status=422,
            detail="empty docx document",
        )
    try:
        return Document(io.BytesIO(data))
    except Exception as exc:
        raise AppError(
            code="VALIDATION_ERROR",
            title="Validation Error",
            status=422,
            detail=f"cannot open docx document: {exc}",
        ) from exc


def _save(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def read_text(data: bytes) -> str:
    """All paragraph texts (tables included), newline-joined."""
    doc = _open(data)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if para.text.strip():
                        parts.append(para.text)
    return "\n".join(parts).strip()


def read_markdown(data: bytes) -> str:
    """Best-effort docx → markdown: headings via style name, bullets kept."""
    doc = _open(data)
    lines: list[str] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading"):
            try:
                level = int(style.replace("heading", "").strip() or "1")
            except ValueError:
                level = 1
            lines.append(f"{'#' * max(1, min(level, 6))} {text}")
            continue
        if style.startswith("list") or text.startswith(("•", "-", "* ")):
            lines.append(f"- {text.lstrip('•-* ').strip()}")
            continue
        lines.append(text)
    return "\n\n".join(lines).strip()


def create_docx(paragraphs: list[ParagraphSpec]) -> bytes:
    """Build a document from declarative paragraphs."""
    if not paragraphs:
        raise AppError(
            code="VALIDATION_ERROR",
            title="Validation Error",
            status=422,
            detail="at least one paragraph required",
        )
    doc = Document()
    for spec in paragraphs:
        _append_spec(doc, spec)
    return _save(doc)


def append_paragraphs(data: bytes, paragraphs: list[ParagraphSpec]) -> bytes:
    """Append paragraphs to an existing document and re-serialize."""
    if not paragraphs:
        raise AppError(
            code="VALIDATION_ERROR",
            title="Validation Error",
            status=422,
            detail="at least one paragraph required",
        )
    doc = _open(data)
    for spec in paragraphs:
        _append_spec(doc, spec)
    return _save(doc)


def set_placeholders(data: bytes, mapping: dict[str, str]) -> bytes:
    """Replace ``{{key}}`` markers (paragraphs + table cells) with values.

    When a paragraph contains a marker the whole paragraph text is rebuilt in
    its first run (formatting of that run wins) — python-docx cannot split a
    marker across runs reliably.
    """
    if not mapping:
        raise AppError(
            code="VALIDATION_ERROR",
            title="Validation Error",
            status=422,
            detail="mapping required",
        )

    def _replace(text: str) -> str:
        def _sub(match: re.Match[str]) -> str:
            key = match.group(1)
            return str(mapping.get(key, match.group(0)))

        return _PLACEHOLDER_RE.sub(_sub, text)

    doc = _open(data)
    for para in doc.paragraphs:
        if "{{" in para.text:
            _rewrite_paragraph(para, _replace(para.text))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if "{{" in para.text:
                        _rewrite_paragraph(para, _replace(para.text))
    return _save(doc)


def _append_spec(doc: Document, spec: ParagraphSpec) -> None:
    if spec.heading is not None:
        level = max(1, min(int(spec.heading), 9))
        doc.add_heading(spec.text, level=level)
        return
    para = doc.add_paragraph()
    if spec.style:
        try:
            para.style = spec.style
        except Exception:
            # Unknown style name — fall back to default text paragraph.
            pass
    run = para.add_run(spec.text)
    run.bold = bool(spec.bold)


def _rewrite_paragraph(para: Any, text: str) -> None:
    """Put ``text`` into the paragraph's first run and blank the rest."""
    runs = para.runs
    if not runs:
        para.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""
