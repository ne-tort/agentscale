"""KP export from run artifacts."""

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from tests.conftest import requires_postgres
from tests.helpers import active_cabinet_headers, register_user


async def _ready_run(client: AsyncClient, headers: dict, cabinet_id: str, suffix: str) -> tuple[str, str]:
    catalog_csv = "part_number,title,price,stock\n910-001793,Mouse,890,in_stock\n"
    cat = await client.post(
        f"/api/v1/cabinets/{cabinet_id}/catalogs/upload",
        headers=headers,
        files={"file": ("distrib.csv", catalog_csv.encode("utf-8"), "text/csv")},
        data={"slug": f"dist-{suffix}", "display_name": "Dist", "trusted_seller": "true"},
    )
    assert cat.status_code == 202, cat.text

    created = await client.post(
        "/api/v1/projects",
        headers=headers,
        json={"slug": f"kp-{suffix}", "display_name": "KP project"},
    )
    project_id = created.json()["id"]
    await client.post(f"/api/v1/projects/{project_id}/open", headers=headers)

    spec = "name,qty,part_number\nMouse Logitech,10,910-001793\n"
    upload = await client.post(
        f"/api/v1/projects/{project_id}/inbox/upload",
        headers=headers,
        files={"file": ("spec.csv", spec.encode("utf-8"), "text/csv")},
        data={"auto_run": "true"},
    )
    run_id = upload.json()["run_id"]
    for phase in ("classify", "search", "rank", "variants", "review"):
        step = await client.post(
            f"/api/v1/projects/{project_id}/runs/{run_id}/advance",
            headers=headers,
            json={"target_phase": phase},
        )
        assert step.status_code == 202, step.text
    return project_id, run_id


@requires_postgres
@pytest.mark.asyncio
async def test_kp_export_uses_catalog_price(client: AsyncClient, unique_suffix: str) -> None:
    reg = await register_user(client, unique_suffix, prefix="kp")
    cabinet_id, headers = await active_cabinet_headers(client, reg, unique_suffix)
    project_id, run_id = await _ready_run(client, headers, cabinet_id, unique_suffix)

    exported = await client.post(
        f"/api/v1/projects/{project_id}/export/kp",
        headers=headers,
        json={"run_id": run_id, "include_alternatives": False},
    )
    assert exported.status_code == 200, exported.text
    body = exported.json()
    assert body["lines_filled"] == 1
    assert body["export_path"].startswith("export/kp-")

    filename = body["export_path"].split("/")[-1]
    download = await client.get(
        f"/api/v1/projects/{project_id}/export/{filename}", headers=headers
    )
    assert download.status_code == 200
    import tempfile
    from pathlib import Path

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp.write(download.content)
        tmp_path = Path(tmp.name)
    book = load_workbook(tmp_path)
    sheet = book.active
    prices = [row[7].value for row in sheet.iter_rows(min_row=2)]
    assert 890 in prices
    assert 891 not in prices


@requires_postgres
@pytest.mark.asyncio
async def test_kp_export_before_review_blocked(client: AsyncClient, unique_suffix: str) -> None:
    reg = await register_user(client, unique_suffix, prefix="kp-early")
    _, headers = await active_cabinet_headers(client, reg, unique_suffix)
    created = await client.post(
        "/api/v1/projects",
        headers=headers,
        json={"slug": f"early-{unique_suffix}", "display_name": "Early"},
    )
    project_id = created.json()["id"]
    await client.post(f"/api/v1/projects/{project_id}/open", headers=headers)
    upload = await client.post(
        f"/api/v1/projects/{project_id}/inbox/upload",
        headers=headers,
        files={"file": ("spec.txt", b"line\n", "text/plain")},
        data={"auto_run": "true"},
    )
    run_id = upload.json()["run_id"]
    exported = await client.post(
        f"/api/v1/projects/{project_id}/export/kp",
        headers=headers,
        json={"run_id": run_id},
    )
    assert exported.status_code == 409
    assert exported.json()["code"] == "PHASE_GUARD_FAILED"


@requires_postgres
@pytest.mark.asyncio
async def test_kp_export_generic_capability_missing(
    client: AsyncClient, unique_suffix: str
) -> None:
    reg = await register_user(client, unique_suffix, prefix="kp-gen")
    _, headers = await active_cabinet_headers(
        client, reg, unique_suffix, profile_id="generic-assistant"
    )
    created = await client.post(
        "/api/v1/projects",
        headers=headers,
        json={"slug": f"gen-{unique_suffix}", "display_name": "Gen"},
    )
    project_id = created.json()["id"]
    await client.post(f"/api/v1/projects/{project_id}/open", headers=headers)
    upload = await client.post(
        f"/api/v1/projects/{project_id}/inbox/upload",
        headers=headers,
        files={"file": ("spec.txt", b"line\n", "text/plain")},
        data={"auto_run": "true"},
    )
    run_id = upload.json()["run_id"]
    for phase in ("classify", "search", "rank", "variants", "review"):
        await client.post(
            f"/api/v1/projects/{project_id}/runs/{run_id}/advance",
            headers=headers,
            json={"target_phase": phase},
        )
    exported = await client.post(
        f"/api/v1/projects/{project_id}/export/kp",
        headers=headers,
        json={"run_id": run_id},
    )
    assert exported.status_code == 403
    assert exported.json()["code"] == "CAPABILITY_MISSING"
