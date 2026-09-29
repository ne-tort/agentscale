"""Unit tests — documents module: local converter + editing helpers."""

from __future__ import annotations

import pytest

from prodavan.application.documents.adapters.local_converter import LocalConverter
from prodavan.application.documents.editing import docx as docx_editing
from prodavan.application.documents.editing import xlsx as xlsx_editing
from prodavan.application.documents.editing.docx_template import render_docx_template
from prodavan.domain.errors import AppError


def _xlsx_bytes(rows: list[list[object]], *, formula_cell: str | None = None) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows:
        ws.append(row)
    if formula_cell:
        ws[formula_cell] = "=SUM(A1:A2)"
    import io as _io

    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes(paragraphs: list[str], *, with_tag: str | None = None) -> bytes:
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if with_tag:
        doc.add_paragraph(with_tag)
    import io as _io

    buf = _io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------- converter


@pytest.mark.asyncio
async def test_local_converter_xlsx_to_csv() -> None:
    data = _xlsx_bytes([["Name", "Price"], ["SSD", 100], ["RAM", 50]])
    out = await LocalConverter().convert(data, "prices.xlsx", "csv")
    assert out.mime == "text/csv"
    assert out.filename == "prices.csv"
    text = out.data.decode("utf-8")
    assert "Name" in text and "SSD" in text


@pytest.mark.asyncio
async def test_local_converter_docx_to_txt() -> None:
    data = _docx_bytes(["Hello documents", "second line"])
    out = await LocalConverter().convert(data, "doc.docx", "txt")
    assert "Hello documents" in out.data.decode("utf-8")
    assert out.filename == "doc.txt"


@pytest.mark.asyncio
async def test_local_converter_identity_passthrough() -> None:
    out = await LocalConverter().convert(b"plain", "notes.txt", "md")
    assert out.data == b"plain"
    assert out.filename == "notes.md"


@pytest.mark.asyncio
async def test_local_converter_unsupported_raises() -> None:
    data = _xlsx_bytes([["a"]])
    with pytest.raises(AppError) as err:
        await LocalConverter().convert(data, "legacy.doc", "pdf")
    assert err.value.code == "CONVERSION_UNAVAILABLE"


@pytest.mark.asyncio
async def test_local_converter_pdf_target_without_gotenberg_message() -> None:
    data = _docx_bytes(["x"])
    with pytest.raises(AppError) as err:
        await LocalConverter().convert(data, "doc.docx", "pdf")
    assert "gotenberg" in (err.value.detail or "").lower()


# ----------------------------------------------------------------------- editing


def test_xlsx_read_rows_and_sheets() -> None:
    data = _xlsx_bytes([["Name", "Qty"], ["SSD", 2]])
    assert xlsx_editing.list_sheets(data) == ["Sheet1"]
    rows = xlsx_editing.read_rows(data)
    assert rows == [{"Name": "SSD", "Qty": 2}]


def test_xlsx_set_cells_roundtrip() -> None:
    data = _xlsx_bytes([["Name", "Note"], ["old", None]])
    out = xlsx_editing.set_cells(data, cells={"A2": "new", "B2": "extra"})
    rows = xlsx_editing.read_rows(out)
    assert rows[0]["Name"] == "new"
    assert rows[0]["Note"] == "extra"


def test_xlsx_fill_template_preserves_formulas() -> None:
    data = _xlsx_bytes([["A", "B"], [1, 2]], formula_cell="C1")
    out = xlsx_editing.fill_template_cells(data, cells={"A2": 10})
    rows = xlsx_editing.read_rows(out)
    assert rows[0]["A"] == 10
    # the untouched formula cell survives the openpyxl round-trip
    from openpyxl import load_workbook

    wb = load_workbook(__import__("io").BytesIO(out))
    assert wb.active["C1"].value == "=SUM(A1:A2)"


def test_xlsx_fill_template_refuses_formula_overwrite() -> None:
    data = _xlsx_bytes([["A"], [1]], formula_cell="A1")
    with pytest.raises(AppError) as err:
        xlsx_editing.fill_template_cells(data, cells={"A1": "value"})
    assert err.value.code == "VALIDATION_ERROR"
    assert "formula" in (err.value.detail or "")


def test_xlsx_create_workbook() -> None:
    out = xlsx_editing.create_workbook(
        sheets=[xlsx_editing.SheetSpec(name="Data", header=["a", "b"], rows=[[1, 2]])]
    )
    rows = xlsx_editing.read_rows(out)
    assert rows == [{"a": 1, "b": 2}]


def test_docx_create_read_and_placeholders() -> None:
    out = docx_editing.create_docx(
        paragraphs=[
            docx_editing.ParagraphSpec(text="Title line", heading=1),
            docx_editing.ParagraphSpec(text="Body {{name}}", bold=False),
        ]
    )
    assert "Title line" in docx_editing.read_text(out)
    filled = docx_editing.set_placeholders(out, {"name": "World"})
    assert "World" in docx_editing.read_text(filled)
    assert "{{name}}" not in docx_editing.read_text(filled)


def test_docx_template_render() -> None:
    data = _docx_bytes(["intro"], with_tag="Hello {{ who }}!")
    out = render_docx_template(data, {"who": "Prodavan"})
    text = docx_editing.read_text(out)
    assert "Hello Prodavan!" in text
    assert "{{" not in text
