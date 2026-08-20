"""Unit tests for spec table cleaning and xlsx ingest (no invented P/N)."""

from pathlib import Path

from openpyxl import Workbook

from prodavan.cabinets.electronics_procurement.pipeline.classify import classify_rows
from prodavan.cabinets.electronics_procurement.pipeline.ingest import ingest_rows
from prodavan.cabinets.electronics_procurement.pipeline.spec_table import clean_extracted_rows


def test_clean_drops_ref_and_prefers_pn_sheet() -> None:
    rows = [
        {
            "sheet": "КП",
            "row": 1,
            "cells": ["№", "Описание", "P/N", "Кол-во"],
            "raw": "header",
        },
        {
            "sheet": "КП",
            "row": 2,
            "cells": ["1", "Mouse Logitech", "", "10"],
            "raw": "no pn",
        },
        {
            "sheet": "Маржа",
            "row": 2,
            "cells": ["1", "Mouse Logitech", "910-001793", "10"],
            "raw": "with pn",
        },
        {
            "sheet": "Спецификация",
            "row": 3,
            "cells": ["2", "#REF!", "#REF!", "1"],
            "raw": "error",
        },
    ]
    cleaned = clean_extracted_rows(rows)
    assert len(cleaned) == 1
    assert "910-001793" in cleaned[0]["cells"]


def test_xlsx_ingest_maps_headers_and_does_not_invent_pn(tmp_path: Path) -> None:
    path = tmp_path / "spec.xlsx"
    wb = Workbook()
    spec = wb.active
    spec.title = "Спецификация"
    spec.append(["№", "Описание", "P/N", "Кол-во"])
    spec.append([1, "Mouse Logitech", "910-001793", 10])
    spec.append([2, "Notebook ACER 16GB NX.FAKE", "", 2])
    dup = wb.create_sheet("КП")
    dup.append(["№", "Описание", "P/N", "Кол-во"])
    dup.append([1, "Mouse Logitech", "910-001793", 10])
    junk = wb.create_sheet("Условия")
    junk.append(["Коммерческое предложение действительно 14 дней"])
    wb.save(path)

    ingest = ingest_rows(path)
    assert ingest["parser"] == "xlsx"
    assert "parser_status" not in ingest
    assert len(ingest["rows"]) == 2

    classified = classify_rows(ingest["rows"])
    mouse = next(i for i in classified["items"] if "910-001793" in i["raw_text"])
    notebook = next(i for i in classified["items"] if "ACER" in i["raw_text"])
    assert mouse["part_number"] == "910-001793"
    assert mouse["qty"] == 10.0
    assert notebook["part_number"] is None


def test_corrupt_xlsx_does_not_invent_rows(tmp_path: Path) -> None:
    path = tmp_path / "spec.xlsx"
    path.write_bytes(b"PK\x03\x04not-a-real-xlsx")
    ingest = ingest_rows(path)
    assert ingest["rows"] == []
    assert ingest["parser_status"] == "error"


def test_xls_stays_not_implemented(tmp_path: Path) -> None:
    path = tmp_path / "legacy.xls"
    path.write_bytes(b"not-ole")
    ingest = ingest_rows(path)
    assert ingest["rows"] == []
    assert ingest["parser_status"] == "not_implemented"
