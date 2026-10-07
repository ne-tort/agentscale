"""Мастер-прайс: отбор поставщиков, OS-запрос, ячейки шаблона, рендер xlsx."""

from __future__ import annotations

import io

import openpyxl
import pytest

from prodavan.application.modules.equipment_master_price import (
    MASTER_PRICE_HEADERS,
    MasterPriceService,
    build_master_price_query,
    doc_to_master_row,
    master_price_row_values,
    master_price_sellers,
    render_master_price_workbook,
)


def _seller(name: str, *, master: bool = True, enabled: bool = True) -> dict:
    return {
        "row_id": f"s_{name}",
        "body": {"name": name, "master_price": master, "is_enabled": enabled},
    }


def test_master_price_template_registered() -> None:
    """Шаблон master_price зарегистрирован и читается из templates/."""
    from prodavan.application.modules.equipment_budget import (
        TEMPLATE_TYPES,
        load_default_template,
    )

    assert TEMPLATE_TYPES["master_price"] == "master-price-template.xlsx"
    data = load_default_template("master_price")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert "Прайс" in wb.sheetnames
    ws = wb["Прайс"]
    assert ws["A1"].value == "Поставщик"
    assert ws["K1"].value == "РРЦ"


def test_master_price_sellers_selection() -> None:
    sellers = [
        _seller("Феррет"),
        _seller("Обычный", master=False),
        _seller("Выключен", enabled=False),
    ]
    assert master_price_sellers(sellers) == ["Феррет"]


def test_build_query_filters_suppliers_and_search() -> None:
    q = build_master_price_query(["Феррет", "Элит"], "")
    assert q["bool"]["filter"] == [{"terms": {"supplier": ["Феррет", "Элит"]}}]
    q = build_master_price_query(["Феррет"], "813661-B21")
    should = q["bool"]["filter"][1]["bool"]["should"]
    assert {"term": {"part_number": "813661-B21"}} in should
    assert q["bool"]["filter"][1]["bool"]["minimum_should_match"] == 1


def test_doc_to_row_and_cells() -> None:
    row = doc_to_master_row(
        {
            "supplier": "Феррет",
            "brand": "HPE",
            "part_number": "813661-B21",
            "title": "Адаптер",
            "in_stock": True,
            "price_num": 100.5,
            "currency": "RUB",
            "src_hash": "h1",
        }
    )
    values = master_price_row_values(row)
    assert values[0] == "Феррет"
    assert values[3] == "813661-B21"
    assert values[5] == "По запросу"  # количеств в каталоге нет
    assert values[7] == "В наличии"
    assert values[8] == 100.5
    assert values[9] == "руб"
    # без цены и под заказ
    poor = doc_to_master_row(
        {"supplier": "X", "in_stock": False, "price_num": None, "currency": "USD"}
    )
    pv = master_price_row_values(poor)
    assert pv[7] == "Под заказ"
    assert pv[8] == "По запросу"
    assert pv[9] == "usd"


def _template_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Прайс"
    for col, header in enumerate(MASTER_PRICE_HEADERS, start=1):
        ws.cell(row=1, column=col, value=header)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_render_workbook_appends_rows_under_header() -> None:
    data = render_master_price_workbook(
        _template_bytes(),
        [
            doc_to_master_row(
                {
                    "supplier": "Феррет",
                    "title": "Т",
                    "in_stock": True,
                    "price_num": 5.0,
                    "currency": "RUB",
                }
            )
        ],
    )
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["Прайс"]
    assert ws["A1"].value == "Поставщик"
    assert ws["A2"].value == "Феррет"
    assert ws["I2"].value == 5.0
    assert ws["J2"].value == "руб"


class _FakeResult:
    def __init__(self, hits, total):
        self.hits = hits
        self.total = total


class _Hit:
    def __init__(self, source, sort):
        self.source = source
        self.sort = sort


class _FakeSearchService:
    def __init__(self, docs):
        self._docs = docs
        self.calls = []

    async def search(self, **kwargs):
        self.calls.append(kwargs)
        from_ = kwargs.get("from_", 0)
        size = kwargs.get("size", 10)
        after = kwargs.get("search_after")
        pool = self._docs
        if after:
            # курсор = [src_hash]: отдаём строго следующие по сортировке
            pool = [d for d in pool if d["src_hash"] > after[-1]]
        elif from_:
            pool = pool[from_:]
        page = pool[:size]
        return _FakeResult([_Hit(d, [d["src_hash"]]) for d in page], len(self._docs))


@pytest.mark.asyncio
async def test_list_page_paging_and_search(monkeypatch) -> None:
    docs = [
        {
            "supplier": "S",
            "title": f"t{i}",
            "src_hash": f"h{i}",
            "part_number": "",
            "in_stock": True,
            "price_num": 1.0,
            "currency": "RUB",
        }
        for i in range(5)
    ]
    fake = _FakeSearchService(docs)
    monkeypatch.setattr(
        "prodavan.core.infra.opensearch_manager.get_search_index_service", lambda: fake
    )
    svc = MasterPriceService(object(), principal=object(), employee=None)

    async def _sellers(cabinet_id):
        return ["S"]

    async def _indexes(cabinet_id):
        return ["equipment__c_row_x"]

    async def _margins(cabinet_id):
        return {}

    monkeypatch.setattr(svc, "_sellers", _sellers)
    monkeypatch.setattr(svc, "_ready_indexes", _indexes)
    monkeypatch.setattr(svc, "_margins", _margins)

    page1 = await svc.list_page(cabinet_id="cab", page=1, page_size=2)
    assert page1["total"] == 5
    assert [r["title"] for r in page1["items"]] == ["t0", "t1"]
    page3 = await svc.list_page(cabinet_id="cab", page=3, page_size=2)
    assert [r["title"] for r in page3["items"]] == ["t4"]
    # глубокая страница за from-лимитом → пусто без ошибки
    deep = await svc.list_page(cabinet_id="cab", page=500, page_size=200)
    assert deep["items"] == []
    # поиск передаётся в запрос
    await svc.list_page(cabinet_id="cab", search="t1")
    assert fake.calls[-1]["query"]["bool"]["filter"][1]["bool"]["should"]


@pytest.mark.asyncio
async def test_iter_rows_cursor_batches(monkeypatch) -> None:
    docs = [
        {
            "supplier": "S",
            "title": f"t{i}",
            "src_hash": f"h{i}",
            "in_stock": True,
            "price_num": 1.0,
            "currency": "RUB",
        }
        for i in range(3)
    ]
    fake = _FakeSearchService(docs)
    monkeypatch.setattr(
        "prodavan.core.infra.opensearch_manager.get_search_index_service", lambda: fake
    )
    svc = MasterPriceService(object(), principal=object(), employee=None)

    async def _sellers(cabinet_id):
        return ["S"]

    async def _indexes(cabinet_id):
        return ["equipment__c_row_x"]

    async def _margins(cabinet_id):
        return {}

    monkeypatch.setattr(svc, "_sellers", _sellers)
    monkeypatch.setattr(svc, "_ready_indexes", _indexes)
    monkeypatch.setattr(svc, "_margins", _margins)
    monkeypatch.setattr(
        "prodavan.application.modules.equipment_master_price._EXPORT_BATCH", 2
    )

    out = [r async for r in svc.iter_rows(cabinet_id="cab")]
    assert [r["title"] for r in out] == ["t0", "t1", "t2"]
    # второй вызов шёл с search_after от последнего хита батча
    assert fake.calls[1].get("search_after") == ["h1"]


def test_apply_margin_rounds_like_commerce() -> None:
    from prodavan.application.modules.equipment_master_price import apply_margin

    assert apply_margin(100.0, 7) == 107.0
    assert apply_margin(100.0, -2) == 98.0
    assert apply_margin(None, 7) is None
    assert apply_margin(10.005, 0) == 10.005
    assert apply_margin(34696.91, 7) == 37125.69


def test_row_values_rrc_cell() -> None:
    row = doc_to_master_row(
        {"supplier": "S", "in_stock": True, "price_num": 10.0, "rrc_num": 15.5, "currency": "RUB"}
    )
    from prodavan.application.modules.equipment_master_price import _apply_margin_to_row

    _apply_margin_to_row(row, {"s": 10.0})
    values = master_price_row_values(row)
    assert values[8] == 11.0  # цена с наценкой
    assert values[10] == 17.05  # РРЦ с наценкой
    # без РРЦ в источнике — «По запросу», как в Commerce
    norrc = doc_to_master_row({"supplier": "S", "in_stock": True, "price_num": 10.0, "rrc_num": None})
    assert master_price_row_values(norrc)[10] == "По запросу"


def test_column_map_projects_rrc() -> None:
    """rrc — каноническое поле: apply_column_map тянет его из источника."""
    from prodavan.application.modules.equipment_catalog_search import apply_column_map

    mapped = apply_column_map(
        {"pn": "X1", "name": "Т", "price": "10", "rrc": "13.5"},
        {"part_number": "pn", "title": "name", "price": "price", "rrc": "rrc"},
    )
    assert mapped["rrc"] == "13.5"
    mapped_no = apply_column_map({"pn": "X1"}, {"part_number": "pn"})
    assert mapped_no["rrc"] == ""
