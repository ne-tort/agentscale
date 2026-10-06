"""Unit tests — equipment budget: fill workbook, sync, export actions."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock
from xml.etree import ElementTree as ET

import pytest

from prodavan.application.modules import equipment_budget as budget
from prodavan.application.modules.module_action_executor import ModuleActionExecutor
from prodavan.domain.errors import AppError
from prodavan.domain.identity import Principal

TEMPLATE = Path(__file__).resolve().parents[2] / "templates" / "kp-template.xlsx"


def _row(**body):
    return dict(body)


# --------------------------------------------------------------- money caches


def test_money_cache_commerce_parity() -> None:
    cache = budget.money_cache(2, 100, 0.22, 0.1)
    assert cache["H"] == 200.0
    assert cache["K"] == 110.0
    assert cache["M"] == 220.0
    assert cache["J"] == round(110 / 1.22, 6)
    assert cache["L"] == round(220 / 1.22, 6)
    assert cache["N"] == 10.0
    assert cache["O"] == 20.0


def test_money_cache_missing_price_is_empty() -> None:
    assert budget.money_cache(1, None, 0.22, 0.1) == {}


# ------------------------------------------------------------ workbook fill


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="template not shipped")
def test_fill_budget_workbook_on_real_template() -> None:
    rows = [
        _row(title="SSD Samsung", part_number="MZ-77Q", qty=2, price_in=100.0,
             vat=0.22, markup=0.1, seller="OCS", comment="В наличии (+)"),
        _row(title="RAM Kingston", part_number="KVR26", qty=1, price_in=50.0,
             vat=0.22, markup=0.1, seller="DNS", comment=""),
    ]
    out = budget.fill_budget_workbook(TEMPLATE.read_bytes(), rows, {})
    assert out and out != TEMPLATE.read_bytes()

    with zipfile.ZipFile(TEMPLATE, "r") as z:
        tpl_shared = z.read("xl/sharedStrings.xml")
        tpl_styles = z.read("xl/styles.xml")
        tpl_sheet2 = z.read("xl/worksheets/sheet2.xml")
    with zipfile.ZipFile(__import__("io").BytesIO(out), "r") as z:
        names = z.namelist()
        assert "xl/sharedStrings.xml" in names
        # styles/sharedStrings are byte-identical: nothing else was touched
        assert z.read("xl/sharedStrings.xml") == tpl_shared
        assert z.read("xl/styles.xml") == tpl_styles
        sheet1 = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
        sheet2 = z.read("xl/worksheets/sheet2.xml")

    # budget values written directly
    assert "MZ-77Q" in sheet1 and "SSD Samsung" in sheet1 and "OCS" in sheet1
    # template formulas preserved (never rewritten)
    assert "PRODUCT(E7,F7)" in sheet1
    assert "PRODUCT(F7,1+I7)" in sheet1
    # cached totals next to the formulas
    assert "220" in sheet1  # M4 total (2 * 110)
    # third sheet (Спецификация) is part of the budget template
    assert "xl/worksheets/sheet3.xml" in names
    # КП sheet: trimmed to 2 items (quoted sheet refs survive in formulas)
    s2 = sheet2.decode("utf-8")
    assert "'Бюджетирование'!D7" in s2
    assert "'Бюджетирование'!D8" in s2
    assert "'Бюджетирование'!D9" not in s2
    assert sheet2 != tpl_sheet2
    # workbook still opens cleanly for openpyxl
    from openpyxl import load_workbook

    wb = load_workbook(__import__("io").BytesIO(out))
    assert "Бюджетирование" in wb.sheetnames and "КП" in wb.sheetnames
    assert "Спецификация" in wb.sheetnames
    assert wb["Бюджетирование"].max_row and wb["Бюджетирование"].max_row >= 100


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="template not shipped")
def test_fill_budget_workbook_empty_rows_placeholders() -> None:
    out = budget.fill_budget_workbook(TEMPLATE.read_bytes(), [], {})
    sheet1 = zipfile.ZipFile(__import__("io").BytesIO(out)).read(
        "xl/worksheets/sheet1.xml"
    ).decode("utf-8")
    # no rows → no D7 value written (placeholder stays untouched)
    assert "Не найден" not in sheet1


# ------------------------------------------------------------- budget actions


def _executor() -> ModuleActionExecutor:
    return ModuleActionExecutor(session=MagicMock())  # type: ignore[arg-type]


async def _invoke(executor, monkeypatch, action, row_id=None):
    async def _load(*args, **kwargs):  # noqa: ANN003
        return action

    monkeypatch.setattr(executor, "_load_action", _load)
    return await executor.invoke(
        cabinet_id="cab_1",
        module_id="mod_equipment",
        action_id=action["id"],
        row_id=row_id,
        principal=Principal(sub="u1", roles=frozenset()),
        employee=None,
    )


@pytest.mark.asyncio
async def test_budget_sync_routes_to_pipeline(monkeypatch) -> None:
    """WAVE7: equipment.budget_sync = полный пайплайн (materialize=True).

    Семантика снапшота (best-оффер, маржа поставщика, ручные правки) покрыта
    tests/unit/test_equipment_offers_pipeline.py.
    """
    executor = _executor()
    calls: list[dict] = []

    async def _pipeline(**kwargs):
        calls.append(kwargs)
        return {"kind": "equipment.pipeline", "materialize": kwargs.get("materialize")}

    monkeypatch.setattr(executor, "_equipment_pipeline", _pipeline)
    out = await _invoke(
        executor,
        monkeypatch,
        {"id": "budget_sync_lines", "kind": "equipment.budget_sync", "params": {}},
    )
    assert out["materialize"] is True
    assert len(calls) == 1
    assert calls[0]["cabinet_id"] == "cab_1"


@pytest.mark.asyncio
async def test_budget_export_saves_document(monkeypatch) -> None:
    executor = _executor()

    async def _rows(**kwargs):  # noqa: ANN003
        if kwargs["table_slug"] == "budget_lines":
            return [
                {"row_id": "b1", "body": {"title": "SSD", "qty": 1, "price_in": 10,
                                          "vat": 0.22, "markup": 0.1, "seller": "OCS"}},
            ]
        return []

    monkeypatch.setattr(executor, "_list_rows_for_scope", _rows)
    monkeypatch.setattr(
        executor, "_resolve_documents_company_id", AsyncMock(return_value="co1")
    )
    ref = {"asset_id": "a", "version_id": "v", "filename": "budget.xlsx"}

    async def _save(_self, data, **kwargs):  # noqa: ANN003
        assert kwargs["filename"] == "budget.xlsx"
        assert kwargs["company_id"] == "co1"
        return ref

    monkeypatch.setattr(
        "prodavan.application.documents.service.DocumentsService.save_document", _save
    )

    out = await _invoke(
        executor,
        monkeypatch,
        {
            "id": "budget_export",
            "kind": "equipment.budget_export",
            "params": {"budget_table": "budget_lines"},
        },
    )
    assert out["kind"] == "equipment.budget_export"
    assert out["file_ref"] == ref
    assert out["rows"] == 1


@pytest.mark.asyncio
async def test_budget_export_requires_rows(monkeypatch) -> None:
    executor = _executor()

    async def _rows(**kwargs):  # noqa: ANN003
        return []

    monkeypatch.setattr(executor, "_list_rows_for_scope", _rows)
    with pytest.raises(AppError) as err:
        await _invoke(
            executor,
            monkeypatch,
            {"id": "budget_export", "kind": "equipment.budget_export", "params": {}},
        )
    assert "no budget lines" in (err.value.detail or "")


@pytest.mark.asyncio
async def test_maybe_auto_equipment_pipeline_routing(monkeypatch) -> None:
    """WAVE7 роутер: found_groups → материализация; lines/offers → пересчёт;
    procurement → procurement_apply; прочие таблицы — ничего."""
    executor = _executor()

    async def _list_actions(*args, **kwargs):  # noqa: ANN003
        return [
            {
                "id": "equipment_pipeline_sync",
                "kind": "equipment.pipeline",
                "enabled": True,
                "params": {
                    "groups_table": "found_groups",
                    "lines_table": "request_lines",
                    "offers_table": "found_offers",
                },
                "trigger": {"on": ["row.created", "row.updated"]},
            },
            {
                "id": "procurement_apply",
                "kind": "equipment.procurement_apply",
                "enabled": True,
                "params": {"procurement_table": "procurement"},
                "trigger": {"on": ["row.updated"]},
            },
        ]

    pipeline_calls: list[dict] = []
    apply_calls: list[str] = []

    async def _pipeline(**kwargs):
        pipeline_calls.append(kwargs)

    async def _apply(**kwargs):
        apply_calls.append(str(kwargs.get("row_id")))

    monkeypatch.setattr(executor, "_list_actions", _list_actions)
    monkeypatch.setattr(executor, "_equipment_pipeline", _pipeline)
    monkeypatch.setattr(executor, "_procurement_apply", _apply)
    principal = Principal(sub="u1", roles=frozenset())

    await executor.maybe_auto_budget_sync(
        cabinet_id="cab_1", module_id="mod_equipment", table_slug="found_groups",
        principal=principal, employee=None,
    )
    await executor.maybe_auto_budget_sync(
        cabinet_id="cab_1", module_id="mod_equipment", table_slug="request_lines",
        principal=principal, employee=None,
    )
    await executor.maybe_auto_budget_sync(
        cabinet_id="cab_1", module_id="mod_equipment", table_slug="procurement",
        row_id="proc_row_1", principal=principal, employee=None,
    )
    await executor.maybe_auto_budget_sync(
        cabinet_id="cab_1", module_id="mod_equipment", table_slug="catalogs",
        principal=principal, employee=None,
    )

    assert [c["materialize"] for c in pipeline_calls] == [True, False]
    assert apply_calls == ["proc_row_1"]


# ------------------------------------------------ standalone КП / spec fill

KP_TEMPLATE = Path(__file__).resolve().parents[2] / "templates" / "commercial-proposal-template.xlsx"
SPEC_TEMPLATE = Path(__file__).resolve().parents[2] / "templates" / "specification-template.xlsx"
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def _kp_cells(out: bytes) -> dict[str, tuple[str | None, bool]]:
    with zipfile.ZipFile(io.BytesIO(out)) as z:
        assert z.testzip() is None
        wb = z.read("xl/workbook.xml").decode("utf-8")
        assert "Бюджетирование" not in wb, "standalone template must not reference the budget sheet"
        part = "xl/worksheets/sheet2.xml"
        if part not in z.namelist():
            part = "xl/worksheets/sheet3.xml"
        sheet = z.read(part)
    root = ET.fromstring(sheet)
    cells: dict[str, tuple[str | None, bool]] = {}
    for row in root.find("m:sheetData", NS).findall("m:row", NS):
        for c in row.findall("m:c", NS):
            v = c.find("m:v", NS)
            is_ = c.find("m:is/m:t", NS)
            f = c.find("m:f", NS)
            val = is_.text if is_ is not None else (v.text if v is not None else None)
            cells[c.get("r")] = (val, f is not None)
    return cells


@pytest.mark.skipif(not KP_TEMPLATE.is_file(), reason="КП template not shipped")
def test_fill_kp_workbook_writes_values_and_trims() -> None:
    rows = [
        _row(title="SSD Samsung 990 Pro 2TB", qty=2, price_in=15000, vat=0.22, markup=0.1),
        _row(title="Патч-корд Vention 3м", qty=10, price_in=250.5, vat=0.22, markup=0.2),
        _row(title="без цены", qty=1, price_in=0, vat=0.22, markup=0.1),
    ]
    cells = _kp_cells(budget.fill_kp_workbook(KP_TEMPLATE.read_bytes(), rows, "commercial_proposal"))
    # values (no formulas) in the item rows (first slot = row 11)
    assert cells["B11"][0] == "SSD Samsung 990 Pro 2TB"
    assert float(cells["E11"][0]) == 2.0
    assert float(cells["G11"][0]) == 16500.0
    assert float(cells["H11"][0]) == 33000.0
    assert not cells["G11"][1] and not cells["H11"][1]
    assert cells["B12"][0] == "Патч-корд Vention 3м"
    assert float(cells["H12"][0]) == 3006.0
    # unpriced rows are skipped; empty tail slots trimmed
    assert "B13" not in cells or cells["B13"][0] != "без цены"
    assert float(cells.get("H13", ("0", False))[0] or 0) != 33000.0


@pytest.mark.skipif(not SPEC_TEMPLATE.is_file(), reason="spec template not shipped")
def test_fill_kp_workbook_spec_variant() -> None:
    rows = [_row(title="Router Mikrotik hEX", qty=1, price_in=5000, vat=0.22, markup=0.1)]
    cells = _kp_cells(
        budget.fill_kp_workbook(SPEC_TEMPLATE.read_bytes(), rows, "specification")
    )
    # spec layout: first slot row 14, columns B/C/D/E
    assert cells["B14"][0] == "Router Mikrotik hEX"
    assert float(cells["C14"][0]) == 1.0
    assert float(cells["D14"][0]) == 5500.0
    assert float(cells["E14"][0]) == 5500.0


def test_fill_kp_workbook_rejects_two_sheet_template() -> None:
    with pytest.raises(AppError) as err:
        budget.fill_kp_workbook(TEMPLATE.read_bytes(), [], "commercial_proposal")
    assert "exactly one sheet" in (err.value.detail or "")


def test_load_default_template_unknown_type() -> None:
    with pytest.raises(AppError):
        budget.load_default_template("nope")


# ------------------------------------------------------- template resolution


@pytest.mark.asyncio
async def test_resolve_export_template_falls_back_without_rows(monkeypatch) -> None:
    executor = _executor()
    fallback = b"FALLBACK"

    async def _rows(**kwargs):  # noqa: ANN003
        return []

    monkeypatch.setattr(executor, "_list_rows_for_scope", _rows)
    out = await executor._resolve_export_template(
        cabinet_id="cab_1",
        template_type="commercial_proposal",
        principal=Principal(sub="u1", roles=frozenset()),
        employee=None,
        fallback=fallback,
    )
    assert out == fallback


@pytest.mark.asyncio
async def test_resolve_export_template_unbound_module_falls_back(monkeypatch) -> None:
    executor = _executor()

    async def _rows(**kwargs):  # noqa: ANN003
        raise AppError(code="NOT_FOUND", title="Not Found", status=404, detail="unbound")

    monkeypatch.setattr(executor, "_list_rows_for_scope", _rows)
    out = await executor._resolve_export_template(
        cabinet_id="cab_1",
        template_type="budget",
        principal=Principal(sub="u1", roles=frozenset()),
        employee=None,
        fallback=b"DEFAULT",
    )
    assert out == b"DEFAULT"


@pytest.mark.asyncio
async def test_resolve_export_template_picks_latest_active(monkeypatch) -> None:
    executor = _executor()

    async def _rows(**kwargs):  # noqa: ANN003
        return [
            {"row_id": "t2", "body": {"template_type": "budget", "active": True,
                                      "file": {"storage_key": "sk2"}}},
            {"row_id": "t1", "body": {"template_type": "budget", "active": False,
                                      "file": {"storage_key": "sk_inactive"}}},
            {"row_id": "t0", "body": {"template_type": "commercial_proposal",
                                      "file": {"storage_key": "sk_other"}}},
        ]

    monkeypatch.setattr(executor, "_list_rows_for_scope", _rows)
    loaded: list[str] = []

    class _Store:
        def get_bytes_sync(self, key: str) -> bytes:
            loaded.append(key)
            return f"XLSX:{key}".encode()

    import prodavan.infrastructure.files.manager as files_mod

    monkeypatch.setattr(files_mod, "ensure_file_store", lambda: _Store())
    out = await executor._resolve_export_template(
        cabinet_id="cab_1",
        template_type="budget",
        principal=Principal(sub="u1", roles=frozenset()),
        employee=None,
        fallback=b"DEFAULT",
    )
    assert out == b"XLSX:sk2"
    assert loaded == ["sk2"]


# ------------------------------------------------------------- КП/спецификация export


@pytest.mark.asyncio
async def test_kp_export_saves_and_converts(monkeypatch) -> None:
    executor = _executor()

    async def _rows(**kwargs):  # noqa: ANN003
        if kwargs["table_slug"] == "budget_lines":
            return [
                {"row_id": "b1", "body": {"title": "SSD", "qty": 1, "price_in": 10,
                                          "vat": 0.22, "markup": 0.1}},
            ]
        return []

    monkeypatch.setattr(executor, "_list_rows_for_scope", _rows)
    monkeypatch.setattr(
        executor, "_resolve_documents_company_id", AsyncMock(return_value="co1")
    )
    saved: dict = {}
    converts: list = []

    async def _save(_self, data, **kwargs):  # noqa: ANN003
        saved["filename"] = kwargs["filename"]
        saved["mime"] = kwargs["mime"]
        saved["head"] = bytes(data)[:5]
        return {"asset_id": "a1", "version_id": "v1", "filename": kwargs["filename"]}

    async def _convert(_self, ref, **kwargs):  # noqa: ANN003
        converts.append(kwargs)
        raise AssertionError("Gotenberg convert больше не используется")

    monkeypatch.setattr(
        "prodavan.application.documents.service.DocumentsService.save_document", _save
    )
    monkeypatch.setattr(
        "prodavan.application.documents.service.DocumentsService.convert", _convert
    )

    out = await _invoke(
        executor,
        monkeypatch,
        {
            "id": "kp_export",
            "kind": "equipment.kp_export",
            "params": {"budget_table": "budget_lines"},
        },
    )
    assert out["kind"] == "equipment.kp_export"
    assert out["file_ref"]["asset_id"] == "a1"
    assert out["rows"] == 1
    assert saved["filename"] == "commercial-proposal.pdf"
    assert saved["mime"] == "application/pdf"
    assert saved["head"] == b"%PDF-"
    assert converts == []


async def test_spec_export_kind(monkeypatch) -> None:
    executor = _executor()

    async def _rows(**kwargs):  # noqa: ANN003
        if kwargs["table_slug"] == "budget_lines":
            return [
                {"row_id": "b1", "body": {"title": "SSD", "qty": 1, "price_in": 10,
                                          "vat": 0.22, "markup": 0.1}},
            ]
        return []

    monkeypatch.setattr(executor, "_list_rows_for_scope", _rows)
    monkeypatch.setattr(
        executor, "_resolve_documents_company_id", AsyncMock(return_value="co1")
    )

    async def _save(_self, data, **kwargs):  # noqa: ANN003
        return {"asset_id": "a1", "version_id": "v1", "filename": kwargs["filename"]}

    async def _convert(_self, ref, **kwargs):  # noqa: ANN003
        return {"asset_id": "a2", "version_id": "v2", "filename": "spec.pdf"}

    monkeypatch.setattr(
        "prodavan.application.documents.service.DocumentsService.save_document", _save
    )
    monkeypatch.setattr(
        "prodavan.application.documents.service.DocumentsService.convert", _convert
    )

    out = await _invoke(
        executor,
        monkeypatch,
        {
            "id": "spec_export",
            "kind": "equipment.spec_export",
            "params": {"budget_table": "budget_lines"},
        },
    )
    assert out["kind"] == "equipment.spec_export"
    # PDF строится из данных: одна запись, без Gotenberg-конвертации
    assert out["file_ref"]["asset_id"] == "a1"


@pytest.mark.asyncio
async def test_kp_export_builds_pdf_directly(monkeypatch) -> None:
    """КП/Спецификация PDF теперь строится из данных (ReportLab) — без
    Gotenberg-конвертации: save_document получает готовый application/pdf."""
    executor = _executor()

    async def _rows(**kwargs):  # noqa: ANN003
        if kwargs["table_slug"] == "budget_lines":
            return [
                {"row_id": "b1", "body": {"title": "SSD 1TB Samsung", "qty": 2,
                                          "price_in": 10000, "vat": 0.22, "markup": 0.1}},
                {"row_id": "b2", "body": {"title": "Коммутатор Cisco 48p", "qty": 1,
                                          "price_in": 250000, "vat": 0.22, "markup": 0.15}},
            ]
        if kwargs["table_slug"] == "document_fields":
            return [{"row_id": "df1", "body": {"contract_number": "2026/1", "city": "г. Москва"}}]
        return []

    monkeypatch.setattr(executor, "_list_rows_for_scope", _rows)
    monkeypatch.setattr(
        executor, "_resolve_documents_company_id", AsyncMock(return_value="co1")
    )
    saved: dict = {}
    ref = {"asset_id": "a", "version_id": "v", "filename": "commercial-proposal.pdf"}

    async def _save(_self, data, **kwargs):  # noqa: ANN003
        saved["data"] = data
        saved.update(kwargs)
        return ref

    monkeypatch.setattr(
        "prodavan.application.documents.service.DocumentsService.save_document", _save
    )

    out = await _invoke(
        executor,
        monkeypatch,
        {
            "id": "kp_export",
            "kind": "equipment.kp_export",
            "params": {"budget_table": "budget_lines", "fields_table": "document_fields"},
        },
    )
    assert out["kind"] == "equipment.kp_export"
    assert out["file_ref"] == ref
    assert saved["filename"] == "commercial-proposal.pdf"
    assert saved["mime"] == "application/pdf"
    assert bytes(saved["data"])[:5] == b"%PDF-"


@pytest.mark.asyncio
async def test_document_fields_body_merges_company_and_deal(monkeypatch) -> None:
    """Реквизиты документов: компания (весь кабинет) — база, сделка (чат)
    перекрывает; пустые строки пропускаются."""
    executor = _executor()
    seen: list[str] = []

    async def _rows(**kwargs):  # noqa: ANN003
        seen.append(kwargs["table_slug"])
        if kwargs["table_slug"] == "document_company_fields":
            return [{"row_id": "c1", "body": {
                "supplier_name": 'ООО "ИТ Взлёт"',
                "city": "г. Москва",
                "app_number": "1",
            }}]
        if kwargs["table_slug"] == "document_fields":
            return [
                {"row_id": "d0", "body": {}},
                {"row_id": "d1", "body": {"city": "г. Тверь", "customer_name": "ООО «Ромашка»"}},
            ]
        return []

    monkeypatch.setattr(executor, "_list_rows_for_scope", _rows)
    merged = await executor._document_fields_body(
        cabinet_id="cab_1",
        module_id="mod_equipment",
        params={},
        principal=Principal(sub="u1", roles=frozenset()),
        employee=None,
        session_id="main",
    )
    assert seen == ["document_company_fields", "document_fields"]
    assert merged == {
        "supplier_name": 'ООО "ИТ Взлёт"',  # из компании
        "city": "г. Тверь",  # сделка перекрыла город компании
        "app_number": "1",
        "customer_name": "ООО «Ромашка»",
    }


@pytest.mark.asyncio
async def test_document_fields_body_survives_missing_company_table(monkeypatch) -> None:
    """Старые кабинеты без company-таблицы: export не падает, берёт сделку."""
    executor = _executor()

    async def _rows(**kwargs):  # noqa: ANN003
        if kwargs["table_slug"] == "document_company_fields":
            raise AppError(code="NOT_FOUND", title="Not Found", status=404)
        return [{"row_id": "d1", "body": {"contract_number": "0001"}}]

    monkeypatch.setattr(executor, "_list_rows_for_scope", _rows)
    merged = await executor._document_fields_body(
        cabinet_id="cab_1",
        module_id="mod_equipment",
        params={},
        principal=Principal(sub="u1", roles=frozenset()),
        employee=None,
        session_id="main",
    )
    assert merged == {"contract_number": "0001"}


@pytest.mark.asyncio
async def test_spec_export_builds_pdf_with_russian_date(monkeypatch) -> None:
    executor = _executor()

    async def _rows(**kwargs):  # noqa: ANN003
        if kwargs["table_slug"] == "budget_lines":
            return [
                {"row_id": "b1", "body": {"title": "Сервер HPE DL380 Gen10 Plus", "qty": 1,
                                          "price_in": 1500000, "vat": 0.22, "markup": 0.12}},
            ]
        return []

    monkeypatch.setattr(executor, "_list_rows_for_scope", _rows)
    monkeypatch.setattr(
        executor, "_resolve_documents_company_id", AsyncMock(return_value="co1")
    )
    saved: dict = {}

    async def _save(_self, data, **kwargs):  # noqa: ANN003
        saved["data"] = data
        saved.update(kwargs)
        return {"asset_id": "a", "version_id": "v"}

    monkeypatch.setattr(
        "prodavan.application.documents.service.DocumentsService.save_document", _save
    )
    out = await _invoke(
        executor,
        monkeypatch,
        {
            "id": "spec_export",
            "kind": "equipment.spec_export",
            "params": {"budget_table": "budget_lines", "fields_table": "document_fields"},
        },
    )
    assert out["kind"] == "equipment.spec_export"
    assert saved["filename"] == "specification.pdf"
    assert bytes(saved["data"])[:5] == b"%PDF-"


def test_rubles_in_words_and_date_ru() -> None:
    from datetime import date

    from prodavan.application.modules.equipment_docs_render import (
        fmt_date_ru,
        rubles_in_words,
    )

    assert rubles_in_words(100000) == "сто тысяч руб. 00 коп."
    assert rubles_in_words(1234.56) == "одна тысяча двести тридцать четыре руб. 56 коп."
    assert rubles_in_words(2000.0) == "две тысячи руб. 00 коп."  # женский род тысяч
    assert fmt_date_ru(date(2026, 10, 5)) == "«05» октября 2026 г."
